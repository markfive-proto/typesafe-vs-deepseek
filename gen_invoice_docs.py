"""Generate 50 synthetic invoices as native PDF/XLSX (born-digital, not scanned photos), with
a ground-truth manifest so extraction accuracy can actually be scored.
"""
import json
import random
from pathlib import Path

from openpyxl import Workbook
from reportlab.lib.pagesizes import letter
from reportlab.lib.units import inch
from reportlab.pdfgen import canvas

random.seed(11)
ROOT = Path(__file__).parent
OUT = ROOT / "invoice_docs"
OUT.mkdir(exist_ok=True)

VENDORS = {
    "goods": ["Northwind Supply Co", "Atlas Hardware", "Fernway Retail", "Cobalt Outfitters"],
    "services": ["Ridgeline Legal Partners", "Vertex Consulting", "Harbor & Co Advisory", "Loom Design Studio"],
    "utilities_subscription": ["BrightRoute Hosting", "CloudSync Inc", "Attic Robotics Cloud", "Nimbus Utilities"],
}
ITEM_NAMES = {
    "goods": ["Steel brackets (box of 50)", "Packing crates", "Safety gloves (case)", "Warehouse shelving unit", "Pallet wrap rolls"],
    "services": ["Contract review (hrs)", "Consulting — Q3 strategy", "Design sprint facilitation", "Litigation support (hrs)", "Onboarding workshop"],
    "utilities_subscription": ["Hosting — Pro plan", "API calls (overage)", "Storage (GB/mo)", "Support plan", "Bandwidth (TB)"],
}


def make_invoice(idx: int) -> dict:
    category = random.choice(list(VENDORS))
    vendor = random.choice(VENDORS[category])
    invoice_number = f"INV-{random.randint(10000, 99999)}"
    date = f"{random.randint(1,12):02d}/{random.randint(1,28):02d}/2026"
    n_items = random.randint(2, 5)
    items = []
    for _ in range(n_items):
        qty = random.randint(1, 12)
        price = round(random.uniform(15, 400), 2)
        items.append({"desc": random.choice(ITEM_NAMES[category]), "qty": qty, "price": price, "amount": round(qty * price, 2)})
    subtotal = round(sum(i["amount"] for i in items), 2)
    tax = round(subtotal * 0.08, 2)
    total = round(subtotal + tax, 2)
    return {
        "vendor": vendor, "invoice_number": invoice_number, "date": date,
        "category": category, "items": items, "subtotal": subtotal, "tax": tax, "total": total,
    }


def write_pdf(inv: dict, path: Path):
    c = canvas.Canvas(str(path), pagesize=letter)
    w, h = letter
    y = h - inch
    c.setFont("Helvetica-Bold", 16)
    c.drawString(inch, y, inv["vendor"])
    c.setFont("Helvetica", 10)
    y -= 24
    c.drawString(inch, y, f"Invoice #: {inv['invoice_number']}")
    y -= 14
    c.drawString(inch, y, f"Date: {inv['date']}")
    y -= 30
    c.setFont("Helvetica-Bold", 10)
    c.drawString(inch, y, "Description")
    c.drawString(4.2 * inch, y, "Qty")
    c.drawString(5.0 * inch, y, "Price")
    c.drawString(6.0 * inch, y, "Amount")
    y -= 6
    c.line(inch, y, 7.3 * inch, y)
    y -= 16
    c.setFont("Helvetica", 10)
    for it in inv["items"]:
        c.drawString(inch, y, it["desc"])
        c.drawString(4.2 * inch, y, str(it["qty"]))
        c.drawString(5.0 * inch, y, f"${it['price']:.2f}")
        c.drawString(6.0 * inch, y, f"${it['amount']:.2f}")
        y -= 16
    y -= 10
    c.line(4.8 * inch, y, 7.3 * inch, y)
    y -= 16
    c.drawString(5.0 * inch, y, "Subtotal:")
    c.drawString(6.0 * inch, y, f"${inv['subtotal']:.2f}")
    y -= 16
    c.drawString(5.0 * inch, y, "Tax (8%):")
    c.drawString(6.0 * inch, y, f"${inv['tax']:.2f}")
    y -= 16
    c.setFont("Helvetica-Bold", 11)
    c.drawString(5.0 * inch, y, "Total:")
    c.drawString(6.0 * inch, y, f"${inv['total']:.2f}")
    c.save()


def write_xlsx(inv: dict, path: Path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Invoice"
    ws["A1"] = inv["vendor"]
    ws["A2"] = f"Invoice #: {inv['invoice_number']}"
    ws["A3"] = f"Date: {inv['date']}"
    ws.append([])
    ws.append(["Description", "Qty", "Price", "Amount"])
    for it in inv["items"]:
        ws.append([it["desc"], it["qty"], it["price"], it["amount"]])
    ws.append([])
    ws.append(["", "", "Subtotal", inv["subtotal"]])
    ws.append(["", "", "Tax (8%)", inv["tax"]])
    ws.append(["", "", "Total", inv["total"]])
    wb.save(str(path))


manifest = {}
for i in range(1, 51):
    inv = make_invoice(i)
    fmt = "pdf" if i % 2 == 1 else "xlsx"
    fname = f"invoice_{i:03d}.{fmt}"
    path = OUT / fname
    if fmt == "pdf":
        write_pdf(inv, path)
    else:
        write_xlsx(inv, path)
    manifest[fname] = inv

(OUT / "ground_truth.json").write_text(json.dumps(manifest, indent=2))
print(f"wrote {len(manifest)} invoices ({sum(1 for f in manifest if f.endswith('.pdf'))} pdf, "
      f"{sum(1 for f in manifest if f.endswith('.xlsx'))} xlsx) to {OUT}")
